#!/usr/bin/env python3
"""Periodic coordination-polyhedron volume sums and geometric union fractions."""
from __future__ import annotations

import argparse
import gzip
import hashlib
import importlib.metadata
import json
import re
import sys
import warnings
import zipfile
from collections import Counter
from datetime import datetime, timezone
from pathlib import Path
import shutil
import time

import numpy as np
import openpyxl
import pandas as pd
from joblib import Parallel, delayed
from pymatgen.core import Structure

ROOT = Path(__file__).resolve().parents[2]
sys.path.insert(0, str(ROOT / "src/analysis"))
sys.path.insert(0, str(ROOT / "src/discovery"))
from polyhedron_geometry import describe_ligand_hull
from periodic_polyhedron_union import periodic_polyhedron_union_fraction
from bonding_vesta import build_vesta_mx_graph, vesta_rule_provenance


def sha(path):
    return hashlib.sha256(Path(path).read_bytes()).hexdigest()


def json_write(path, value):
    Path(path).write_text(json.dumps(value, ensure_ascii=False, indent=2, allow_nan=False), encoding="utf-8")


def cell_geometry(structure, material_id, cell_kind, config):
    lattice = np.asarray(structure.lattice.matrix)
    row = {"material_id": material_id, "cell_kind": cell_kind,
           "cell_volume_A3": float(structure.volume), "n_sites": len(structure),
           "lattice_matrix_A": lattice.tolist(), "lattice_abc_A": list(structure.lattice.abc),
           "lattice_angles_deg": list(structure.lattice.angles), "status": "pending"}
    records, hulls, geometry_payload = [], [], []
    try:
        graph = build_vesta_mx_graph(structure)
    except (KeyError, ValueError) as exc:
        row.update(status="bonding_unavailable", error_reason=str(exc))
        return row, records, hulls, geometry_payload
    row.update(center_element=graph.center_symbol, halogen_element=graph.halogen_symbol,
               cutoff_A=float(graph.cutoff_A), n_centers=len(graph.center_indices), n_bonds=len(graph.bonds))
    if not graph.bonds:
        row.update(status="no_vesta_bonds", error_reason="VESTA has no M-X bonds; geometric fraction is unavailable")
        return row, records, hulls, geometry_payload
    by_center = graph.bonds_by_center()
    failed = False
    for index in graph.center_indices:
        bonds = by_center.get(index, ())
        frac_points = np.asarray([np.asarray(structure[b.halogen_index].frac_coords) + np.asarray(b.image) for b in bonds], dtype=float).reshape(-1, 3)
        vectors = (frac_points - np.asarray(structure[index].frac_coords)) @ lattice
        description = describe_ligand_hull(vectors, center_cart=np.zeros(3),
            abs_tol_A=config["hull_rank_abs_tol_A"], rel_tol=config["hull_rank_rel_tol"])
        eligible = len(bonds) >= config["minimum_cn_for_polyhedron"]
        full_3d = eligible and description["geometry_status"] == "full_3d"
        if eligible and description["geometry_status"] == "hull_error":
            failed = True
        assigned_volume = float(description["volume_A3"]) if full_3d else (None if eligible and description["geometry_status"] == "hull_error" else 0.0)
        record = {"material_id": material_id, "cell_kind": cell_kind, "center_index": int(index),
            "CN": len(bonds), "eligible_CN_ge3": eligible, **description,
            "assigned_geometric_volume_A3": assigned_volume}
        records.append(record)
        geometry_payload.append({**record, "center_fractional_coords": np.asarray(structure[index].frac_coords).tolist(),
            "ligand_fractional_coords_with_images": frac_points.tolist(), "ligand_displacement_vectors_A": vectors.tolist(),
            "ligand_site_image_keys": [[int(b.halogen_index), *map(int, b.image)] for b in bonds]})
        if full_3d:
            hulls.append(frac_points)
    eligible_records = [r for r in records if r["eligible_CN_ge3"]]
    full_records = [r for r in eligible_records if r["geometry_status"] == "full_3d"]
    row.update(n_eligible_shells=len(eligible_records), n_3d_polyhedra=len(full_records),
        n_lower_rank_eligible_shells=sum(r["geometry_status"] != "full_3d" and r["geometry_status"] != "hull_error" for r in eligible_records),
        n_centers_below_min_CN=sum(not r["eligible_CN_ge3"] for r in records),
        n_hull_errors=sum(r["geometry_status"] == "hull_error" for r in eligible_records),
        n_near_planar_3d=sum(bool(r["near_planar"]) for r in full_records),
        n_center_outside_3d_hull=sum(r["center_inside_hull"] is False for r in full_records),
        fraction_centers_with_3d_hull=len(full_records) / len(records),
        valid_3d_volume_sum_A3=float(sum(r["assigned_geometric_volume_A3"] for r in full_records)))
    if failed:
        row.update(status="hull_error", error_reason="At least one eligible shell failed geometry; partial sum is not reported as complete")
    else:
        row["polyhedron_volume_sum_A3"] = row["valid_3d_volume_sum_A3"]
        row["sum_volume_fraction"] = row["polyhedron_volume_sum_A3"] / row["cell_volume_A3"]
        row["status"] = "all_eligible_shells_3d" if len(full_records) == len(eligible_records) and full_records else (
            "includes_lower_rank_shells" if eligible_records else "empty_polyhedron_set_CN_less_than_3")
        if eligible_records and not full_records:
            row["status"] = "only_lower_rank_shells_zero_3d_volume"
    return row, records, hulls, geometry_payload


def process_one(source, config):
    material_id = source["material_id"]
    result = {"material_id": material_id, "formula": source["formula"], "cif_file": source["cif_file"]}
    path = ROOT / "data/candidates/materials_project/cif" / source["cif_file"]
    digest = sha(path)
    if digest != source["cif_sha256"]:
        raise ValueError(f"Source CIF hash changed: {material_id}")
    result["cif_sha256"] = digest
    with warnings.catch_warnings(record=True) as caught:
        warnings.simplefilter("always")
        try:
            raw = Structure.from_file(path)
            standard = raw.get_primitive_structure(tolerance=config["primitive_tolerance_A"], use_site_props=False).get_reduced_structure(reduction_algo="niggli")
        except Exception as exc:
            result.update(status="cif_or_standardization_error", error_reason=str(exc))
            return result, [], [], []
    result["parse_warnings"] = "; ".join(dict.fromkeys(str(w.message) for w in caught))
    raw_row, raw_centers, _, raw_payload = cell_geometry(raw, material_id, "raw_CIF", config)
    std_row, std_centers, hulls, std_payload = cell_geometry(standard, material_id, "primitive_niggli", config)
    for prefix, row in [("raw", raw_row), ("standard", std_row)]:
        result.update({f"{prefix}__{k}": v for k, v in row.items() if k not in ["material_id", "cell_kind"]})
    result["status"] = std_row["status"]
    cell_rows = [raw_row, std_row]
    if "sum_volume_fraction" not in std_row:
        result["error_reason"] = std_row.get("error_reason", "Geometry unavailable")
        return result, raw_centers + std_centers, cell_rows, raw_payload + std_payload
    seed = (config["base_seed"] + int(material_id.split("-")[-1])) % (2 ** 32)
    union = periodic_polyhedron_union_fraction(hulls, sobol_power=config["sobol_power_initial"],
        n_replicates=config["sobol_replicates"], seed=seed, chunk_size=config["sobol_chunk_size"])
    if union.get("error_reason"):
        result.update(union__status="error", union_error_reason=union["error_reason"])
        return result, raw_centers + std_centers, cell_rows, raw_payload + std_payload
    upper = min(1.0, std_row["sum_volume_fraction"])
    raw_mean = union["fraction_mean"]
    needs_refinement = (union["fraction_replicate_range"] > config["replicate_range_refinement_threshold"]
        or raw_mean > upper + config["sum_upper_bound_excess_refinement_threshold"]
        or (raw_mean == 0 and upper > 0))
    result["union_refined"] = needs_refinement
    result["union_initial_replicate_range"] = union["fraction_replicate_range"]
    if needs_refinement:
        union = periodic_polyhedron_union_fraction(hulls, sobol_power=config["sobol_power_refined"],
            n_replicates=config["sobol_replicates"], seed=seed, chunk_size=config["sobol_chunk_size"])
    if union.get("error_reason"):
        result.update(union__status="error", union_error_reason=union["error_reason"])
        return result, raw_centers + std_centers, cell_rows, raw_payload + std_payload
    result["union_second_refined"] = (union["fraction_replicate_range"] > config["replicate_range_refinement_threshold"]
        or union["fraction_mean"] > upper + config["sum_upper_bound_excess_refinement_threshold"]
        or (union["fraction_mean"] == 0 and upper > 0))
    if result["union_second_refined"]:
        result["union_before_final_refinement"] = union
        union = periodic_polyhedron_union_fraction(hulls, sobol_power=config["sobol_power_final"],
            n_replicates=config["sobol_replicates"], seed=seed, chunk_size=config["sobol_chunk_size"])
        if union.get("error_reason"):
            result.update(union__status="error", union_error_reason=union["error_reason"])
            return result, raw_centers + std_centers, cell_rows, raw_payload + std_payload
    if not np.isclose(union["sum_volume_fraction"], std_row["sum_volume_fraction"], atol=1e-9, rtol=1e-7):
        raise ValueError(f"Cartesian/fractional hull volumes disagree for {material_id}")
    result.update({f"union__{k}": v for k, v in union.items()})
    result["union_volume_A3_estimate"] = union["fraction_mean"] * standard.volume
    result["union_fraction_upper_bound_from_sum"] = upper
    result["sum_minus_union_fraction_estimate"] = std_row["sum_volume_fraction"] - union["fraction_mean"]
    result["union_convergence_review"] = (union["fraction_replicate_range"] > config["replicate_range_refinement_threshold"]
        or union["fraction_mean"] > upper + config["sum_upper_bound_excess_refinement_threshold"]
        or (union["fraction_mean"] == 0 and upper > 0))
    result["union_all_samples_inside"] = bool(hulls and union["fraction_mean"] == 1)
    if "sum_volume_fraction" in raw_row:
        result["raw_standard_sum_fraction_difference"] = raw_row["sum_volume_fraction"] - std_row["sum_volume_fraction"]
    return result, raw_centers + std_centers, cell_rows, raw_payload + std_payload


def original_ids(path):
    book = openpyxl.load_workbook(path, read_only=True, data_only=True)
    ids = []
    for sheet in book.worksheets:
        for values in sheet.iter_rows(values_only=True):
            ids.extend(str(x).strip() for x in values if x is not None and re.fullmatch(r"mp-\d+", str(x).strip()))
    book.close()
    if len(ids) != 17 or len(set(ids)) != 17:
        raise ValueError("Unexpected original workbook ID set")
    return ids


STATUS_ZH = {
    "all_eligible_shells_3d": "所有CN≥3配位壳均形成三维多面体",
    "includes_lower_rank_shells": "含三维多面体及低秩配位壳",
    "only_lower_rank_shells_zero_3d_volume": "仅低秩配位壳，三维几何体积为0",
    "empty_polyhedron_set_CN_less_than_3": "所有中心CN<3，定义的多面体集合为空",
    "bonding_unavailable": "无可用VESTA元素对成键参数",
    "no_vesta_bonds": "VESTA规则未识别出M–X键，体积留空",
    "hull_error": "凸包计算失败，完整体积留空",
    "cif_or_standardization_error": "CIF解析或标准化失败",
}
FEATURES = {"vesta__dim": "维度", "vesta__Xcn": "st1", "vesta__Xsh": "st2",
            "vesta__Pcn": "st3", "chem__delta_chi_x_m": "电负性差"}
VOLUME_COLUMNS = {
    "standard__cell_volume_A3": "标准晶胞体积_A3",
    "standard__polyhedron_volume_sum_A3": "多面体体积之和_A3",
    "standard__sum_volume_fraction": "体积求和比值",
    "union_volume_A3_estimate": "并集体积估计_A3",
    "union__fraction_mean": "去重叠占据比例_估计",
    "union__fraction_standard_error": "占据比例_数值标准误",
    "union__fraction_replicate_range": "占据比例_四次估计极差",
    "union__fraction_replicates": "占据比例_四次估计值",
    "union__n_points_per_replicate": "每次Sobol采样点数",
    "union__n_points_evaluated": "最终总采样点数",
    "union__method": "并集积分方法",
    "union_refined": "是否加密采样",
    "union_second_refined": "是否加密至2的18次方",
    "union_convergence_review": "是否需数值收敛审查",
    "standard__n_centers": "中心原子数",
    "standard__n_eligible_shells": "CN≥3配位壳数",
    "standard__n_3d_polyhedra": "三维多面体数",
    "standard__n_lower_rank_eligible_shells": "CN≥3低秩配位壳数",
    "standard__n_centers_below_min_CN": "CN<3中心数",
    "standard__fraction_centers_with_3d_hull": "形成三维多面体的中心比例",
    "standard__n_near_planar_3d": "接近平面的三维多面体数",
    "standard__n_center_outside_3d_hull": "中心位于卤素凸包外的三维多面体数",
    "standard__cutoff_A": "VESTA截断距离_A",
    "raw__cell_volume_A3": "原CIF晶胞体积_A3",
    "raw__sum_volume_fraction": "原CIF体积求和比值",
    "raw_standard_sum_fraction_difference": "原CIF减标准胞_求和比值差",
    "union_fraction_upper_bound_from_sum": "并集占比理论上限",
    "sum_minus_union_fraction_estimate": "求和减并集_含积分误差",
}


def tabular(frame):
    """Encode array-valued diagnostics without altering scalar precision."""
    result = frame.copy()
    for col in result.select_dtypes(include="object"):
        result[col] = result[col].map(lambda x: json.dumps(x, ensure_ascii=False, allow_nan=False)
                                     if isinstance(x, (list, dict, tuple)) else x)
    return result


def export_results(outdir, config, pool, ids, computed, provenance):
    materials = pd.DataFrame([entry[0] for entry in computed])
    centers = pd.DataFrame([row for entry in computed for row in entry[1]])
    cells = pd.DataFrame([row for entry in computed for row in entry[2]])
    if len(materials) != 933 or not materials.material_id.is_unique or set(materials.material_id) != set(pool.material_id):
        raise ValueError("Computed structure coverage differs from the 933 sources")
    order = ids + pool.loc[~pool.material_id.isin(ids), "material_id"].tolist()
    materials = materials.set_index("material_id").loc[order].reset_index()
    source = pool.set_index("material_id").loc[order].reset_index()
    for col in [*FEATURES, "theoretical", "topology__unbonded_halogen_fraction", "feature_status"]:
        materials[col] = source[col]
    materials["original_workbook_17"] = materials.material_id.isin(ids)
    if not materials.theoretical.eq(False).all():
        raise ValueError("Nonexperimental source found")

    complete = materials["standard__sum_volume_fraction"].notna()
    union_complete = materials["union__fraction_mean"].notna()
    if not complete.equals(source.feature_status.eq("ok")):
        raise ValueError("Geometry availability differs from frozen VESTA feature availability")
    if not complete.equals(union_complete):
        raise ValueError("Some supported geometric sets lack a union result; review before export")
    if not materials.loc[union_complete, "union__fraction_mean"].between(0, 1).all():
        raise ValueError("Union estimate lies outside [0,1]")
    if materials.loc[complete, "standard__sum_volume_fraction"].lt(0).any():
        raise ValueError("Negative geometric volume")
    grouped = centers.loc[centers.cell_kind.eq("primitive_niggli")].groupby("material_id")["assigned_geometric_volume_A3"].sum()
    expected = materials.loc[complete].set_index("material_id")["standard__polyhedron_volume_sum_A3"]
    if not np.allclose(grouped.loc[expected.index], expected, atol=1e-9, rtol=1e-10):
        raise ValueError("Per-center sum and per-structure volume disagree")

    main = materials[["formula", "material_id"]].rename(columns={"formula": "材料", "material_id": "MP_ID"}).copy()
    for key, label in VOLUME_COLUMNS.items():
        main[label] = materials[key]
    main.loc[materials["union__method"].eq("exact_empty_set"), "每次Sobol采样点数"] = 0
    main.insert(5, "体积求和百分比", 100 * materials["standard__sum_volume_fraction"])
    main.insert(8, "去重叠占据百分比_估计", 100 * materials["union__fraction_mean"])
    main["几何状态"] = materials.status.map(STATUS_ZH).fillna(materials.status)
    main["几何状态代码"] = materials.status
    main["体积结果可用"] = np.where(complete, "是", "否")
    main["原表17结构"] = np.where(materials.original_workbook_17, "是", "否")
    main["Experimentally Observed"] = "Yes"
    main["未成键卤素比例"] = materials["topology__unbonded_halogen_fraction"]
    main["成键审查提示"] = np.where(main["未成键卤素比例"].gt(0), "部分卤素在VESTA规则下未成键", "")
    main["缺失原因"] = materials.get("error_reason", pd.Series("", index=materials.index)).fillna("")
    for key, label in FEATURES.items():
        main[label] = materials[key]
    main["CIF文件"] = materials.cif_file
    main["CIF_SHA256"] = materials.cif_sha256
    main = tabular(main)
    original = main.iloc[:17].copy()
    remaining = main.iloc[17:].copy()
    zero = complete & materials["standard__sum_volume_fraction"].eq(0)
    positive = complete & materials["standard__sum_volume_fraction"].gt(0)
    review = materials.union_convergence_review.eq(True)
    cell_diff = materials.raw_standard_sum_fraction_difference.dropna().abs()

    definitions = pd.DataFrame([
        ("范围", "本地MP快照中933个theoretical=False结构；原Excel17个MP结构加其余916个。保留全部MP ID，不按体积、组成或结构重新去重。"),
        ("实验记录", "Experimentally Observed=Yes指二元母体在MP有实验观察记录；所用MP CIF坐标不因此成为原始实测坐标，也不表示衍生电解质已经证实粘弹性。"),
        ("成键定义", "冻结VESTA-2019元素对距离阈值；严格d<cutoff，保留周期镜像，无其他算法回退。沿用此前st1/st2/st3定义。"),
        ("多面体顶点", "一个中心周围所有成键卤素原子核（含必要周期镜像）的凸包；中心原子不作为附加顶点。每个晶胞内的中心只计一次。"),
        ("晶胞", "原CIF经0.01 Å容差求原胞，再作Niggli约化，与已有五特征计算一致；另列原CIF晶胞求和结果供比较。"),
        ("体积求和比值", "ΣV(完整局域配位凸包)/V(晶胞)，无量纲；共享体积会重复计数，因此允许大于1。百分比列为比值×100。"),
        ("去重叠占据比例_估计", "所有三维凸包及其周期复制体在一个晶胞内的并集体积/V(晶胞)，在0至1之间。使用分数坐标域[0,1)^3上的独立扰乱Sobol积分。"),
        ("数值积分", "初始4×2^14点；四次极差>0.002、超过min(1,求和比值)+0.0005或非零凸包全无命中时加密至4×2^16点，仍触发则加密至4×2^18点。保留加密后未收敛提示和各次估计。"),
        ("数值标准误与极差", "仅反映独立Sobol扰乱的积分波动，不是严格误差上界或材料物理不确定性；零波动或全命中也不证明精确收敛。"),
        ("求和减并集", "含积分误差的有符号差值；轻微负值可来自有限采样，未强行截为0。"),
        ("低秩与零值", "有效非空成键图中，CN<3不属于既有多面体集合；CN≥3但卤素顶点仿射秩<3的配位壳三维体积定义为0，不人为加厚。若不存在三维壳则总量为0，状态单列。"),
        ("秩与几何覆盖率", "SVD阈值max(1e-8 Å,1e-7×最大奇异值)；近乎平面标志为最小/最大奇异值<1e-4。均记录逐中心秩、CN、凸包体积和中心是否位于凸包内。"),
        ("空白", "无VESTA元素对参数、规则下全无M–X键或计算失败时留空；不能将不可计算条目视为0。"),
        ("配位覆盖", "形成三维多面体的中心比例及未成键卤素比例辅助审查；混合低秩壳、低配位中心及中心在凸包外的结构保留并标记。"),
        ("物理含义", "这是配位多面体的几何占据描述符，不是原子堆积率；1减占据比例不能直接解释为可用自由体积，也不能单独证明粘弹性。"),
        ("维度/st1/st2/st3/电负性差", "沿用VESTA字段：维度=vesta__dim，st1=Xcn，st2=Xsh(共享卤素数)，st3=Pcn(相邻多面体数)，Δχ=χ卤素−χ中心。"),
        ("详细文件", "materials_detailed.csv保留所有计算诊断；逐中心及晶胞CSV含原CIF与标准胞；geometry_checkpoint.json.gz保留完整周期配位坐标以便复核。"),
    ], columns=["项目", "定义与说明"])
    tables = {"原表17个结构": original, "其余916个结构": remaining, "全部933个结构": main,
              "三维体积大于0": main.loc[positive], "三维几何体积为0": main.loc[zero],
              "体积不可计算": main.loc[~complete], "积分收敛待审查": main.loc[review],
              "逐中心几何": tabular(centers), "晶胞对照": tabular(cells), "字段定义与说明": definitions}
    for name, frame in tables.items():
        frame.to_csv(outdir / f"{name}.csv", index=False, encoding="utf-8-sig")
    tabular(materials).to_csv(outdir / "materials_detailed.csv", index=False, encoding="utf-8-sig")
    excel = outdir / "MP实验记录933个结构_多面体体积占比.xlsx"
    with pd.ExcelWriter(excel, engine="openpyxl") as writer:
        for name, frame in tables.items():
            frame.to_excel(writer, sheet_name=name, index=False)
            sheet = writer.sheets[name]
            sheet.freeze_panes = "C2"
            sheet.auto_filter.ref = sheet.dimensions
            for cell in sheet[1]:
                cell.font = openpyxl.styles.Font(bold=True, color="FFFFFF")
                cell.fill = openpyxl.styles.PatternFill("solid", fgColor="225577")
            for index, col in enumerate(frame.columns, 1):
                width = 60 if col in ["定义与说明", "缺失原因", "CIF_SHA256"] else 24
                if col == "几何状态":
                    width = 48
                sheet.column_dimensions[openpyxl.utils.get_column_letter(index)].width = width
    numeric_columns = ["标准晶胞体积_A3", "多面体体积之和_A3", "体积求和比值", "去重叠占据比例_估计", *FEATURES.values()]
    for name in ["原表17个结构", "其余916个结构", "全部933个结构"]:
        back = pd.read_excel(excel, sheet_name=name)
        if back.MP_ID.tolist() != tables[name].MP_ID.tolist() or not np.allclose(
            back[numeric_columns].astype(float), tables[name][numeric_columns].astype(float),
            equal_nan=True, rtol=1e-12, atol=1e-12):
            raise ValueError(f"Excel roundtrip failed: {name}")
    counts = {"all_structures": len(main), "original_workbook": len(original), "remaining": len(remaining),
        "volume_available": int(complete.sum()), "positive_3d_volume": int(positive.sum()),
        "zero_3d_geometric_volume": int(zero.sum()), "volume_unavailable": int((~complete).sum()),
        "refined_union_integration": int(materials.union_refined.eq(True).sum()),
        "second_refined_union_integration": int(materials.union_second_refined.eq(True).sum()),
        "integration_review": int(review.sum()), "geometry_status": dict(Counter(materials.status)),
        "partial_unbonded_halogen_structures": int(main["未成键卤素比例"].gt(0).sum()),
        "structures_with_center_outside_hull": int(materials["standard__n_center_outside_3d_hull"].gt(0).sum()),
        "structures_with_near_planar_3d": int(materials["standard__n_near_planar_3d"].gt(0).sum())}
    validation = {"exported_at_utc": datetime.now(timezone.utc).isoformat(), "counts": counts,
        "all_933_source_cif_hashes_verified": True, "all_theoretical_false": True,
        "original_17_exact_ids_verified": True, "per_center_volume_sums_verified": True,
        "availability_matches_frozen_feature_support": True, "excel_roundtrip_verified": True,
        "max_raw_standard_sum_fraction_absolute_difference": float(cell_diff.max()),
        "raw_standard_difference_over_1e_6_count": int(cell_diff.gt(1e-6).sum()),
        "raw_standard_difference_over_1e_6_ids": materials.loc[materials.raw_standard_sum_fraction_difference.abs().gt(1e-6), "material_id"].tolist(),
        "max_union_replicate_range_after_refinement": float(materials["union__fraction_replicate_range"].max()),
        "union_minus_analytical_upper_bound_max": float((materials["union__fraction_mean"] - materials.union_fraction_upper_bound_from_sum).max()),
        "nonzero_hull_no_sample_hit_ids": materials.loc[positive & materials["union__fraction_mean"].eq(0), "material_id"].tolist(),
        "integration_review_ids": materials.loc[review, "material_id"].tolist(),
        "provenance": provenance}
    json_write(outdir / "validation_report.json", validation)
    json_write(outdir / "calculation_config.json", config)
    readme = (
        "# MP实验观察结构：配位多面体体积占比\n\n"
        f"本地MP快照共933个结构，原Excel17个+其余916个；{counts['volume_available']}个体积结果可用，"
        f"其中{counts['positive_3d_volume']}个三维体积>0，{counts['zero_3d_geometric_volume']}个按定义三维体积为0；"
        f"{counts['volume_unavailable']}个因成键规则不可计算而留空。\n\n"
        "主要列：`体积求和比值`=所有完整局域配位凸包体积之和/晶胞体积（可>1）；"
        "`去重叠占据比例_估计`=周期并集体积/晶胞体积（0至1，数值积分估计）。百分比列均为对应比值×100。\n\n"
        "多面体采用冻结VESTA规则下成键的卤素原子核为顶点，保留跨晶胞边界的完整配位，中心原子不另加为顶点。"
        "主结果使用与此前五特征一致的原胞+Niggli标准化。原CIF结果、逐中心CN/秩/体积及完整周期顶点均保留。\n\n"
        "CN≥3但平面/线状配位壳的三维体积为0；有效成键图中所有中心CN<3时定义集合为空，其体积也为0。"
        "这些零值有明确几何状态，与没有成键参数或完全没有识别出M–X键的缺失值分开。\n\n"
        f"并集初始4×2^14个独立扰乱Sobol点，{counts['refined_union_integration']}个结构加密至4×2^16点，"
        f"其中{counts['second_refined_union_integration']}个进一步加密至4×2^18点；"
        f"最终{counts['integration_review']}个仍有收敛审查提示。各次估计、极差和标准误均导出，仅表征积分波动，不能作为严格误差上界。"
        "求和减并集的轻微负值可能来自积分误差，未截断。\n\n"
        "本指标是配位多面体几何占据，不是原子堆积率、可用自由体积或粘弹性的直接判据。"
        "MP实验记录属于母体，MP结构坐标不因此等同于原始实验精修坐标。\n\n"
        "Excel包含17、916、933分表，以及零体积、缺失、收敛提示、逐中心、晶胞对照和字段定义。"
        "materials_detailed.csv保留机器可读字段；geometry_checkpoint.json.gz记录全部周期顶点；"
        "validation_report.json含来源及脚本SHA256、成键表哈希和检查结果。\n\n"
        "在项目根目录复现（使用新的输出目录）：\n\n```bash\n"
        "MKL_THREADING_LAYER=GNU OPENBLAS_NUM_THREADS=1 OMP_NUM_THREADS=1 MKL_NUM_THREADS=1 \\\n"
        "python src/analysis/compute_experimental_polyhedron_volume.py --outdir results/experimental_polyhedron_volume_rerun\n```\n"
    )
    (outdir / "README.md").write_text(readme, encoding="utf-8")
    reproduction = outdir / "reproduce"
    reproduction.mkdir(exist_ok=True)
    for relative in ["src/analysis/compute_experimental_polyhedron_volume.py", "src/analysis/polyhedron_geometry.py",
                     "src/analysis/periodic_polyhedron_union.py", "src/discovery/bonding_vesta.py",
                     "configs/experimental_polyhedron_volume_v1.json", "tests/test_polyhedron_geometry.py",
                     "tests/test_periodic_polyhedron_union.py"]:
        target = reproduction / relative
        target.parent.mkdir(parents=True, exist_ok=True)
        shutil.copy2(ROOT / relative, target)
    archive = outdir / "MP实验结构933个_多面体体积占比.zip"
    with zipfile.ZipFile(archive, "w", zipfile.ZIP_DEFLATED) as zf:
        for path in sorted(outdir.rglob("*")):
            if path.is_file() and path != archive:
                zf.write(path, path.relative_to(outdir))
    with zipfile.ZipFile(archive) as zf:
        if zf.testzip() is not None:
            raise ValueError("ZIP integrity failure")
    print(json.dumps({"excel": str(excel), "zip": str(archive), "counts": counts}, ensure_ascii=False), flush=True)


def main():
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("--config", type=Path, default=ROOT / "configs/experimental_polyhedron_volume_v1.json")
    parser.add_argument("--outdir", type=Path, default=ROOT / "results/experimental_polyhedron_volume_v1")
    parser.add_argument("--export-only", action="store_true", help="Re-export a verified checkpoint without recomputing geometry")
    args = parser.parse_args()
    config = json.loads(args.config.read_text(encoding="utf-8"))
    outdir = args.outdir.resolve()
    pool_path = ROOT / config["pool_table"]
    workbook = ROOT / config["original_workbook"]
    pool = pd.read_csv(pool_path)
    ids = original_ids(workbook)
    metadata = pd.read_csv(ROOT / "data/candidates/materials_project/metadata.csv")
    expected = set(metadata.loc[metadata.theoretical.eq(False), "material_id"])
    if len(pool) != 933 or not pool.material_id.is_unique or not pool.theoretical.eq(False).all() or set(pool.material_id) != expected or not set(ids).issubset(set(pool.material_id)):
        raise ValueError("Unexpected experimental-pool coverage")
    rule = vesta_rule_provenance(verify=True)
    sources = [pool_path, workbook, args.config.resolve(), ROOT / config["feature_definition_config"],
               Path(__file__).resolve(), ROOT / "src/analysis/polyhedron_geometry.py",
               ROOT / "src/analysis/periodic_polyhedron_union.py", ROOT / "src/discovery/bonding_vesta.py",
               ROOT / "data/candidates/materials_project/metadata.csv"]
    provenance = {"vesta_bonding": rule, "source_sha256": {str(p.relative_to(ROOT)): sha(p) for p in sources},
        "versions": {name: importlib.metadata.version(name) for name in ["pymatgen", "numpy", "scipy", "pandas", "openpyxl", "joblib"]},
        "python_version": sys.version,
        "cif_sha256_by_mp_id": dict(zip(pool.material_id, pool.cif_sha256))}
    checkpoint = outdir / "geometry_checkpoint.json.gz"
    if args.export_only:
        with gzip.open(checkpoint, "rt", encoding="utf-8") as handle:
            saved = json.load(handle)
        if saved["config"] != config:
            raise ValueError("Checkpoint configuration differs")
        for relative, digest in saved["provenance"]["source_sha256"].items():
            if relative != str(Path(__file__).resolve().relative_to(ROOT)) and sha(ROOT / relative) != digest:
                raise ValueError(f"Checkpoint source has changed: {relative}")
        for row in pool.to_dict("records"):
            if sha(ROOT / "data/candidates/materials_project/cif" / row["cif_file"]) != row["cif_sha256"]:
                raise ValueError(f"CIF changed: {row['material_id']}")
        computed = saved["computed"]
        provenance["computation_provenance"] = saved["provenance"]
    else:
        outdir.mkdir(parents=True, exist_ok=False)
        started = time.monotonic()
        computed = []
        generator = Parallel(n_jobs=config["n_jobs"], return_as="generator", batch_size="auto")(
            delayed(process_one)(row, config) for row in pool.to_dict("records"))
        for index, entry in enumerate(generator, 1):
            computed.append(entry)
            if index % 50 == 0 or index == len(pool):
                print(f"Computed {index}/{len(pool)} structures in {time.monotonic() - started:.1f}s", flush=True)
        with gzip.open(checkpoint, "wt", encoding="utf-8") as handle:
            json.dump({"config": config, "provenance": provenance, "computed": computed}, handle, ensure_ascii=False, allow_nan=False)
    export_results(outdir, config, pool, ids, computed, provenance)


if __name__ == "__main__":
    main()
